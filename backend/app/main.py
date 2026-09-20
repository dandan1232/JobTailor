import io
import re
from contextlib import asynccontextmanager
from pathlib import Path

from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from pypdf import PdfReader

from backend.app.ai import AIAnalysisError, ai_is_configured, request_ai_analysis


MAX_RESUME_BYTES = 5 * 1024 * 1024
ALLOWED_RESUME_SUFFIXES = {".pdf", ".docx"}
SKILLS = (
    "Python", "FastAPI", "Django", "PostgreSQL", "MySQL", "Redis", "Docker",
    "Kubernetes", "AWS", "React", "TypeScript", "Java", "Spring Boot", "Vue",
    "Node.js", "微服务", "高并发", "团队协作",
)


class AnalyzeRequest(BaseModel):
    resume_text: str = Field(min_length=40, max_length=30_000)
    job_description: str = Field(min_length=30, max_length=20_000)


class Evidence(BaseModel):
    skill: str
    evidence: str


class Gap(BaseModel):
    skill: str
    suggestion: str


class Dimension(BaseModel):
    name: str
    score: int
    note: str


class Requirement(BaseModel):
    label: str
    status: str
    evidence: str


class Revision(BaseModel):
    id: str
    priority: str
    category: str
    title: str
    original: str
    revised: str
    reason: str


class AnalyzeResponse(BaseModel):
    analysis_mode: str
    score: int
    verdict: str
    dimensions: list[Dimension]
    requirements: list[Requirement]
    matched: list[Evidence]
    gaps: list[Gap]
    revisions: list[Revision]


class ResumeExtractResponse(BaseModel):
    filename: str
    characters: int
    text: str


def contains_skill(text: str, skill: str) -> bool:
    return skill.casefold() in text.casefold()


def fragments(text: str) -> list[str]:
    return [part.strip(" •\t-") for part in re.split(r"[\n。；]", text) if part.strip()]


def evidence_for(resume_text: str, skill: str) -> str:
    return next(
        (part for part in fragments(resume_text) if contains_skill(part, skill)),
        "简历中已提及该项能力",
    )


def extract_target_role(job_description: str) -> str:
    for line in fragments(job_description)[:4]:
        cleaned = re.sub(r"^(岗位|职位|招聘职位|Job Title)\s*[:：]\s*", "", line, flags=re.I)
        if cleaned != line and 2 <= len(cleaned) <= 40:
            return cleaned
    return "目标岗位"


def best_project_line(resume_text: str) -> str:
    lines = fragments(resume_text)
    quantified = [line for line in lines if any(char.isdigit() for char in line)]
    return (quantified or lines or ["未识别到可用项目描述"])[0]


def build_revisions(
    resume_text: str,
    role: str,
    matched_skills: list[str],
    gap_skills: list[str],
) -> list[Revision]:
    project_line = best_project_line(resume_text)
    skill_summary = "、".join(matched_skills[:4]) or "相关技术"
    revisions = [
        Revision(
            id="summary-focus",
            priority="高",
            category="个人摘要",
            title="让开头直接回应目标岗位",
            original="简历开头缺少针对当前岗位的能力定位。",
            revised=f"面向{role}岗位，具备{skill_summary}等相关经验，能够独立推进服务开发、问题定位与交付。",
            reason="招聘方通常先扫读顶部摘要。先说明岗位方向和已有证据，能更快建立相关性。",
        ),
        Revision(
            id="project-result",
            priority="中",
            category="项目经历",
            title="把项目描述改成动作与结果",
            original=project_line,
            revised=f"围绕业务目标说明你的具体动作、使用的技术和可验证结果。原始证据：{project_line}",
            reason="只罗列职责很难判断贡献大小。动作、技术选择和结果能让经历更可追问。",
        ),
    ]

    if gap_skills:
        missing = "、".join(gap_skills[:3])
        revisions.insert(
            1,
            Revision(
                id="skill-gap",
                priority="高",
                category="技能缺口",
                title=f"确认 {missing} 是否有真实证据",
                original=f"JD 要求 {missing}，当前简历没有找到直接证据。",
                revised=f"如确实使用过 {missing}，补充对应项目、你的职责和结果；如果没有，不要为了匹配度强行加入。",
                reason="缺少证据的关键词容易在面试追问时失真。这里应补真实经历，而不是补漂亮话。",
            ),
        )

    has_quantified_result = any(
        "%" in line or re.search(r"\d+\s*(秒|人|次|万|小时|天)", line)
        for line in fragments(resume_text)
    )
    if not has_quantified_result:
        revisions.append(
            Revision(
                id="quantify-impact",
                priority="中",
                category="成果表达",
                title="补充可核实的结果",
                original="多数经历没有结果数据或明确影响。",
                revised="在有真实记录时补充性能、效率、规模或交付结果；没有准确数字时可用定性结果，不要编造。",
                reason="结果证据能区分“参与过”和“真正产生过影响”。",
            ),
        )

    return revisions


def analyze_resume(payload: AnalyzeRequest) -> AnalyzeResponse:
    role = extract_target_role(payload.job_description)
    requirements = [skill for skill in SKILLS if contains_skill(payload.job_description, skill)]
    matched_skills = [skill for skill in requirements if contains_skill(payload.resume_text, skill)]
    gap_skills = [skill for skill in requirements if skill not in matched_skills]

    coverage = len(matched_skills) / max(len(requirements), 1)
    has_numbers = any(char.isdigit() for char in payload.resume_text)
    has_bullets = payload.resume_text.count("\n") >= 3
    keyword_score = round(coverage * 100)
    evidence_score = min(58 + (18 if has_numbers else 0) + len(matched_skills) * 4, 96)
    relevance_score = min(45 + round(coverage * 45) + (8 if role != "目标岗位" else 0), 96)
    clarity_score = 84 if has_bullets else 66
    score = round(
        keyword_score * 0.35
        + evidence_score * 0.3
        + relevance_score * 0.2
        + clarity_score * 0.15
    )

    dimensions = [
        Dimension(name="关键词覆盖", score=keyword_score, note=f"命中 {len(matched_skills)}/{len(requirements)} 项识别要求"),
        Dimension(name="经历证据", score=evidence_score, note="检查技能是否有项目与结果支撑"),
        Dimension(name="岗位相关性", score=relevance_score, note=f"围绕{role}判断经历关联度"),
        Dimension(name="表达清晰度", score=clarity_score, note="检查结构、动作和结果是否易读"),
    ]
    hard_requirements = [
        Requirement(
            label=skill,
            status="pass" if skill in matched_skills else "missing",
            evidence=evidence_for(payload.resume_text, skill) if skill in matched_skills else "简历中暂未找到直接证据",
        )
        for skill in requirements[:6]
    ]
    matched = [Evidence(skill=skill, evidence=evidence_for(payload.resume_text, skill)) for skill in matched_skills[:6]]
    gaps = [
        Gap(skill=skill, suggestion=f"如有 {skill} 的真实经历，请补充项目、职责和结果；否则保留为缺口。")
        for skill in gap_skills[:4]
    ]
    verdict = "匹配基础较好，优先补齐高优先级证据后再投递。" if score >= 70 else "存在明显要求缺口，建议先核实真实经历并调整表达。"

    return AnalyzeResponse(
        analysis_mode="local",
        score=score,
        verdict=verdict,
        dimensions=dimensions,
        requirements=hard_requirements,
        matched=matched,
        gaps=gaps,
        revisions=build_revisions(payload.resume_text, role, matched_skills, gap_skills),
    )


def extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_docx(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_rows = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    return "\n".join(paragraphs + table_rows)


@asynccontextmanager
async def lifespan(_: FastAPI):
    yield


app = FastAPI(title="JobTailor API", version="0.2.0", lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000", "http://127.0.0.1:3000",
        "http://localhost:3001", "http://127.0.0.1:3001",
    ],
    allow_credentials=False,
    allow_methods=["GET", "POST"],
    allow_headers=["Content-Type"],
)


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok", "analysis_mode": "ai" if ai_is_configured() else "local"}


@app.post("/api/resume/extract", response_model=ResumeExtractResponse)
async def extract_resume(file: UploadFile = File(...)) -> ResumeExtractResponse:
    filename = file.filename or "resume"
    suffix = Path(filename).suffix.casefold()
    if suffix not in ALLOWED_RESUME_SUFFIXES:
        raise HTTPException(status_code=415, detail="仅支持 PDF 和 DOCX 简历")

    content = await file.read(MAX_RESUME_BYTES + 1)
    if len(content) > MAX_RESUME_BYTES:
        raise HTTPException(status_code=413, detail="简历文件不能超过 5MB")

    try:
        text = extract_pdf(content) if suffix == ".pdf" else extract_docx(content)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="无法读取该简历，请确认文件未损坏或加密") from exc

    normalized = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(normalized) < 40:
        raise HTTPException(status_code=422, detail="未提取到足够文字，扫描版 PDF 暂不支持")
    normalized = normalized[:30_000]

    return ResumeExtractResponse(filename=filename, characters=len(normalized), text=normalized)


@app.post("/api/analyze", response_model=AnalyzeResponse)
async def analyze(payload: AnalyzeRequest) -> AnalyzeResponse:
    try:
        ai_result = await request_ai_analysis(payload.resume_text, payload.job_description)
        if ai_result is not None:
            return AnalyzeResponse.model_validate({"analysis_mode": "ai", **ai_result})
    except (AIAnalysisError, ValueError) as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    return analyze_resume(payload)


@app.post("/api/match", response_model=AnalyzeResponse, deprecated=True)
def match_job(payload: AnalyzeRequest) -> AnalyzeResponse:
    return analyze_resume(payload)
