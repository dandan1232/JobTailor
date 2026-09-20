import io

from docx import Document
from fastapi.testclient import TestClient

from backend.app.main import app


client = TestClient(app)


def test_health_reports_local_mode(monkeypatch) -> None:
    monkeypatch.delenv("AI_BASE_URL", raising=False)
    monkeypatch.delenv("AI_API_KEY", raising=False)
    monkeypatch.delenv("AI_MODEL", raising=False)

    response = client.get("/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok", "analysis_mode": "local"}


def test_extracts_docx_resume_text() -> None:
    document = Document()
    document.add_heading("陈默 - Python 后端工程师", level=1)
    document.add_paragraph("使用 FastAPI 开发订单服务，并通过 Redis 将响应时间降低 42%。")
    stream = io.BytesIO()
    document.save(stream)

    response = client.post(
        "/api/resume/extract",
        files={
            "file": (
                "resume.docx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["filename"] == "resume.docx"
    assert "FastAPI" in body["text"]
    assert body["characters"] > 40


def test_rejects_unsupported_resume_type() -> None:
    response = client.post(
        "/api/resume/extract",
        files={"file": ("resume.txt", b"plain text resume", "text/plain")},
    )

    assert response.status_code == 415
    assert response.json()["detail"] == "仅支持 PDF 和 DOCX 简历"


def test_analyze_returns_dimensions_requirements_and_revisions(monkeypatch) -> None:
    monkeypatch.delenv("AI_BASE_URL", raising=False)
    monkeypatch.delenv("AI_API_KEY", raising=False)
    monkeypatch.delenv("AI_MODEL", raising=False)
    response = client.post(
        "/api/analyze",
        json={
            "resume_text": "使用 Python 和 FastAPI 开发订单服务，引入 Redis 后响应时间降低 42%。\n通过 Docker 完成交付。",
            "job_description": "岗位：Python 后端工程师\n要求熟悉 Python、FastAPI、Redis、Docker 和 Kubernetes，负责后端服务开发。",
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["analysis_mode"] == "local"
    assert len(body["dimensions"]) == 4
    assert {item["skill"] for item in body["matched"]} == {"Python", "FastAPI", "Redis", "Docker"}
    assert {item["skill"] for item in body["gaps"]} == {"Kubernetes"}
    assert any(item["id"] == "skill-gap" for item in body["revisions"])
